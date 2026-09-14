#!/usr/bin/env python3
"""Audit a workbook against the modelling conventions in SKILL.md.

Usage:
    python audit.py <file.xlsx> [--strict] [--sheet NAME]

Checks, in order of importance:

    formulas              derived numbers are formulas, not pasted results
    formula_literal       a formula body that is only a typed constant (`=0`)
    cached_error          a formula whose cached value is an error (#DIV/0!, #REF!, or the
                          bare #ERROR! IronCalc returns for a construct it cannot evaluate)
    formula_hardcode      a numeric literal typed inside a formula body (`=B4*0.42`)
    functions             dynamic-array or 365-only functions; functions_prefix for a
                          name that needs `_xlfn.`; functions_network and
                          functions_volatile for those families
    references            every cross-sheet reference names a real sheet;
                          references_external for a link into another workbook, whether
                          the cell names it or a defined name does it for the cell
    colors_input          a typed number not styled blue; colors_formula_blue for a
                          formula styled as an input
    formula_font_color    a formula font off the code; red is for external links only
    provenance            every hardcoded input carries a cell comment naming its source
    solved_unchecked      a `Solved:` input that no Checks verdict depends on, so nothing
                          fails when its target moves
    solved_comment        a `Solved:` comment that names no target cell or no re-solve
                          command
    solved_color          a `Solved:` cell that is not a blue input
    solved_formula        a `Solved:` comment on a formula; a solved value is typed, and
                          a number a formula can produce is not one
    formats_percent       a fraction typed on `General`, which is a rate with no format
    formula_number_format a formula that evaluates to a number left on `General`
    totals                a row labelled Total/Subtotal/Sum/Net holds a typed number;
                          totals_formula when its formula is neither a SUM, an explicit
                          addition or subtraction, nor a bare link
    frozen_row            a formula row whose text does not change across periods
    formula_family        one formula in a row that does not share the row's shape
    formula_density       a calculation sheet whose numbers are typed rather than computed
    cover_static          a first sheet with no formulas, so its tiles go stale
    sensitivity_centre    a grid centre that does not reproduce the output it varies
    sensitivity_unverified a grid whose centre could not be checked at all
    checks_missing        a sensitivity grid with no Checks sheet asserting its direction
    checks_rollup         a Checks sheet with no verdict formula or no FAIL roll-up
    hidden_sheet          a sheet the reader never sees
    iterative_calc        iterative calculation on, which means a circular reference
    layout_widths         most used columns still on the default width
    layout_freeze         no frozen header row on a long sheet
    layout_gridlines      gridlines visible (info)

Every check reports counts plus up to 25 example cells so the agent can fix
the workbook without reopening it cell by cell. Findings are advisory; with
--strict the exit code is 1 when any "fail"-level finding exists, and
formula_number_format is raised from warn to fail.

"Input" means a numeric constant in a sheet that also contains formulas.
"Derived" means a numeric cell whose row or column already contains a formula
(the same line item across periods, or the same period down a schedule).
Label columns and header rows are excluded, and a sheet with no formulas at all
is raw data or pure inputs: it is never scored for density and its colours are
not judged, though a blue cell on it still needs its comment. Formula density
reads the same way on a calculation sheet: a blue declared input and a period
header row are both set aside, and every other typed number counts against the
sheet that shows it.

`--sheet` audits one sheet and is an error when the workbook has no sheet by that
name, rather than a pass over nothing; the workbook-level checks (cover, Checks
roll-up, iterative calculation) read the whole file either way.

Thresholds and definitions the checks apply:

    formula_density       counted per sheet from 20 numeric cells up; warns below
                          0.80 of them being formulas, fails below 0.50
    period header row     row 1, and any row holding no formula whose numbers are all
                          whole years from 1990 to 2100 (with a second number or an
                          `FY2026E`/`Q1 2026` label beside them, so one stray integer is
                          not an axis). Its numbers are what the sheet is ruled with, so
                          they are not counted against its density.
    formula_family        a row of 4 or more formulas where one shape covers 65% of
                          them; 1 to 5 cells off that shape are reported, and a run
                          of one shape at either end of the row is exempt (historical
                          links before the forecast, a terminal column after it)
    frozen_row            3 or more adjacent period columns holding identical text
    structural numbers    0, 1, -1, 2, 10, 12, 100, 360, 365, 1000, 1000000, any
                          integer up to 12, an exponent after `^`, and every literal
                          in a verdict formula (one that returns OK/FAIL/WARN/PASS/
                          CHECK/ERROR), since a tolerance or a band edge is what such a
                          formula holds; on a sheet named Checks a literal directly after
                          a comparison is exempt and any other literal is not. A percent
                          literal elsewhere is never structural: `15%` in a formula is an
                          assumption.
    sensitivity grid      a filled rectangle at least 3 by 3 of one formula shape,
                          with numeric headers in the column left of it and the row
                          above it, where every cell reads the two headers that cross
                          on it. Nothing about the sheet names is consulted.
    sensitivity centre    the middle cell of an odd-sided grid, checked against the
                          output the workbook designates: the cell paired with the centre
                          in a Checks row that carries a verdict, else the cell a
                          `=centre-output` difference names anywhere, else the cell a data
                          table's formula cell reads, where the two axis headers meet. It
                          has to agree to within 0.5%, and a grid that designates nothing
                          is sensitivity_unverified rather than a fail against whichever
                          named cell happens to sit nearest the centre.
    solved value          a blue input whose comment begins `Solved:`; the one derived
                          number allowed as a typed value. Its comment names the target
                          cell it was solved against, the spot cell it was matched to and
                          the command that re-solves it (`Solved: <what> at which <target>
                          equals <spot>; re-solve with <command>`), and a Checks verdict
                          must depend on it, directly or through the formulas, ranges and
                          defined names that read it.
    cover                 a workbook of three or more sheets opens on a sheet that links
                          to the outputs it reports; cover_static fires when the first
                          visible sheet holds no such link
    examples              up to 25 cells listed per finding
"""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

DYNAMIC_ARRAY = {"XLOOKUP", "XMATCH", "FILTER", "SORT", "SORTBY", "UNIQUE", "SEQUENCE", "LET", "LAMBDA", "RANDARRAY"}
NEEDS_XLFN = {"TEXTJOIN", "CONCAT", "IFS", "SWITCH", "MAXIFS", "MINIFS", "IFNA", "FORECAST.LINEAR", "STDEV.S", "STDEV.P", "VAR.S", "VAR.P", "AGGREGATE"}
VOLATILE = {"INDIRECT", "OFFSET", "NOW", "TODAY", "RAND", "RANDBETWEEN"}
NETWORK = {"WEBSERVICE", "RTD", "FILTERXML", "ENCODEURL"}
# Excel serialises a worksheet dynamic-array function with both prefixes, `_xlfn._xlws.FILTER`,
# and the name after them cannot start the match itself, so the whole prefix is matched here
# or FILTER and SORT read as no function at all.
FUNC_RE = re.compile(r"(?<![A-Za-z0-9_.])(_xlfn\.(?:_xlws\.)?)?([A-Za-z][A-Za-z0-9._]*)\s*\(", re.I)
SHEET_REF_RE = re.compile(r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.]+))!")
TOTAL_LABEL = re.compile(r"\b(total|subtotal|sum|net income|net debt|net)\b", re.I)
# A cell reference whose column is not anchored with $ (sheet-qualified or not).
RELATIVE_REF = re.compile(r"(?<![$A-Za-z0-9_])[A-Z]{1,3}\$?\d+(?![A-Za-z0-9_(])")
# A formula whose whole body is a typed number, with or without wrapping parentheses.
LITERAL_FORMULA = re.compile(r"^=\s*\(*\s*[+-]?(?:\d+\.?\d*|\.\d+)\s*\)*\s*$")
HEX_COLOR = re.compile(r"^(?:[0-9A-Fa-f]{6}|[0-9A-Fa-f]{8})$")
EXAMPLES = 25

BLUE = {"FF0000FF", "000000FF", "0000FF"}
BLACK = {"FF000000", "00000000", "000000", None}
GREEN = {"FF008000", "00008000", "008000", "FF00B050", "0000B050", "00B050"}
RED = {"FFFF0000", "00FF0000", "FF0000"}
TEXT_FUNCS = {"TEXT", "CONCAT", "CONCATENATE", "TEXTJOIN", "LEFT", "RIGHT", "MID", "UPPER", "LOWER", "PROPER", "TRIM", "SUBSTITUTE", "REPLACE", "REPT", "CHAR", "DOLLAR", "FIXED"}

MASK = "\x01"
STRING_LITERAL = re.compile(r'"[^"]*"')
# A sheet or workbook qualifier, a cell or range reference, and a bare name, in that order:
# every one of them can carry digits that are not a typed number.
SHEET_QUALIFIER = re.compile(r"(?:'(?:[^']|'')*'|\[[^\]]*\]|[A-Za-z_][A-Za-z0-9_.]*)!")
# A cell or range, a whole-row range (13:13) or a whole-column range (A:A). The lookbehind
# keeps the E6 of 1.5E6 from reading as a cell.
CELL_RANGE = re.compile(r"(?<![A-Za-z0-9_.])(?:\$?[A-Za-z]{1,3}\$?\d+(?::\$?[A-Za-z]{1,3}\$?\d+)?|\$?\d+:\$?\d+|\$?[A-Za-z]{1,3}:\$?[A-Za-z]{1,3})(?![A-Za-z0-9_(])")
# A structured reference, whole: `Table1[Revenue 2025]`, `Table1[[#Data],[Revenue 2025]]`,
# `[@Cost]`. A column heading is prose, so a year or a size inside the brackets is part of
# a name; masking only the table and the text tokens leaves that year behind as a literal.
STRUCTURED_REF = re.compile(r"(?:[A-Za-z_][A-Za-z0-9_.]*)?\[(?:[^\[\]]|\[[^\[\]]*\])*\]")
BARE_NAME = re.compile(r"(?<![A-Za-z0-9_.])[A-Za-z_][A-Za-z0-9_.]*")
NUMBER_LITERAL = re.compile(r"(?<![A-Za-z0-9_.$])((?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?)(%?)")
# Unit conversions and day counts read as structure, not as an assumption someone typed.
STRUCTURAL_NUMBERS = {0, 1, -1, 2, 10, 12, 100, 360, 365, 1000, 1000000}
STRUCTURAL_INT_MAX = 12
VERDICT_TOKEN = re.compile(r'"\s*(?:OK|FAIL|WARN|PASS|CHECK|ERROR)\s*"', re.I)
CHECKS_SHEET = re.compile(r"^checks?$", re.I)
ERROR_VALUE = re.compile(r"^#(?:N/A|[A-Z0-9/]+[!?])$")
SOLVED_COMMENT = re.compile(r"^\s*solved\s*:", re.I)
# A cell address, sheet-qualified or bare: the comment has to say which cell the value
# was solved against, and the word "target" on its own says nothing a check can use.
SOLVED_TARGET = re.compile(r"(?<![A-Za-z0-9_])(?:'[^']+'!|[A-Za-z_][A-Za-z0-9_.]*!)?\$?[A-Za-z]{1,3}\$?\d+(?![A-Za-z0-9_])")
SOLVED_RESOLVE = re.compile(r"re-?solve\s+with\s+\S+|goal ?seek", re.I)
RANGE_EXPAND_MAX = 4000
FAMILY_TOKEN = re.compile(r"(?:'[^']*'|\[[^\]]*\]|[A-Za-z_][A-Za-z0-9_.]*)!|(?<![A-Za-z0-9_.])\$?[A-Za-z]{1,3}\$?\d+(?![A-Za-z0-9_(])")
CELL_TOKEN = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")
FAMILY_MIN_CELLS = 4
FAMILY_DOMINANT = 0.65
FAMILY_MAX_OUTLIERS = 5

# Formula density. Our own model sheets run 0.96 to 1.00 once declared inputs are set
# aside, so a sheet where one number in five was typed is worth a look and a sheet where
# most of them were typed is a table of results wearing a model's layout.
DENSITY_MIN_CELLS = 20
DENSITY_FAIL = 0.50
DENSITY_WARN = 0.80

# A period header rules a sheet with its columns, so the years in it are not numbers the
# model failed to work out. The window is wide enough for a long history and a long
# forecast and narrow enough that a headcount, a share count or a price is not read as a year.
HEADER_YEAR_MIN = 1990
HEADER_YEAR_MAX = 2100
PERIOD_LABEL = re.compile(r"^\s*(?:(?:FY|CY)\s*'?\d{2,4}|[1-4]Q\s*'?\d{2,4}|Q[1-4](?:\s*'?\d{2,4})?|(?:19|20)\d{2})\s*[EAP]?\s*$", re.I)

# Sensitivity grids: a rectangle of one formula shape whose cells each read the two axis
# headers that cross on them. Shape alone, so a grid on the model's own sheet, which is
# where the DCF convention puts it, is recognised as readily as one that reads an Inputs tab.
GRID_MIN_SIDE = 3
GRID_CENTRE_TOLERANCE = 0.005
# A formula whose whole body is one reference, which is what a cover tile is.
LINK_FORMULA = re.compile(r"^=\s*(?:'[^']+'|[A-Za-z_][A-Za-z0-9_. ]*)!\$?[A-Za-z]{1,3}\$?\d+\s*$")


def rgb(font) -> str | None:
    """Normalise a font colour to its hex digits, or None when it is not a plain RGB one.

    openpyxl hands back its own type-error sentence from `color.rgb` for theme and
    indexed colours, so an isinstance check alone would read that sentence as a
    colour; anything that is not six or eight hex digits counts as the default.
    """
    color = getattr(font, "color", None)
    if color is None or not isinstance(color.rgb, str) or not HEX_COLOR.match(color.rgb):
        return None
    return color.rgb.upper()


def evaluates_to_number(text: str, cached) -> bool:
    """Whether a formula produces a number, from its cached value when the file has one.

    A workbook saved by openpyxl carries no cached values, so the fallback reads the
    formula: string literals, concatenation and the text functions mark it as text and
    everything else is assumed numeric. An empty literal is not one of them, because
    `=IF(B2="","",B2*2)` is a number with a blank for the period that has no driver yet;
    reading its fallback as text dropped every guarded formula out of the density count
    and out of the format and colour checks. A body holding nothing else is still text.
    """
    if cached is not None:
        return is_number(cached)
    body = text.replace('""', "")
    if not body.strip().strip("="):
        return False
    if '"' in body or "&" in body:
        return False
    return not any(name.upper() in TEXT_FUNCS for _, name in FUNC_RE.findall(body))


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def period_header_rows(cells: list, formula_rows: set[int]) -> set[int]:
    """Rows that rule a sheet with its periods rather than carrying a line item.

    Row 1 always, then any formula-free row whose every number is a year: a model repeats
    its period axis above each block, and counting those years as numbers the sheet failed
    to work out is what read a full model as a table of pasted results.
    """
    by_row: dict[int, list] = {}
    for c in cells:
        by_row.setdefault(c.row, []).append(c)
    headers = {1}
    for row, line in by_row.items():
        if row == 1 or row in formula_rows:
            continue
        numbers = [c.value for c in line if is_number(c.value)]
        labels = [c for c in line if isinstance(c.value, str) and PERIOD_LABEL.match(c.value)]
        if not numbers or len(numbers) + len(labels) < 2:
            continue  # one stray integer is a value someone typed, not an axis
        if all(float(v).is_integer() and HEADER_YEAR_MIN <= v <= HEADER_YEAR_MAX for v in numbers):
            headers.add(row)
    return headers


def mask_formula(text: str) -> str:
    """Blank out everything in a formula that is not a typed number, keeping the offsets.

    Masking is the whole check: `$B$12`, `A1:G3`, `Table1[Revenue 2025]` and `LOG10` all carry
    digits, so a number pattern run over the raw text reads references, table headings and
    function names as hardcoded inputs.
    """
    masked = STRING_LITERAL.sub(lambda m: MASK * len(m.group()), text)
    for pattern in (STRUCTURED_REF, SHEET_QUALIFIER, CELL_RANGE, BARE_NAME):
        masked = pattern.sub(lambda m: MASK * len(m.group()), masked)
    return masked


def hardcoded_literals(text: str, sheet: str = "") -> list[str]:
    """Numbers typed into a formula body that stand for an assumption rather than structure."""
    found: list[str] = []
    masked = mask_formula(text)
    # A verdict formula is exempt whole, and only that: it has to render a verdict, or sit
    # on the Checks sheet. Any quoted string used to qualify, so a business threshold walked
    # out of `=IF(B5<0.15,B5,"")` on the strength of a blank-string fallback.
    if VERDICT_TOKEN.search(text):
        return found  # a tolerance, a band edge or a bound: what a verdict formula holds
    on_checks = bool(CHECKS_SHEET.match(sheet))
    for m in NUMBER_LITERAL.finditer(masked):
        prefix = masked[: m.start()].rstrip()
        if prefix.endswith("^"):
            continue  # an exponent is structure: a square, a cube, an n-th root
        if on_checks and len(prefix) > 1 and prefix[-1] in "<>=":
            continue  # a bound compared against on the Checks sheet; `=B5*0.42` there is still a hardcode
        raw, percent = m.group(1), m.group(2)
        if percent:
            found.append(raw + "%")
            continue
        value = float(raw)
        if value in STRUCTURAL_NUMBERS or (value.is_integer() and abs(value) <= STRUCTURAL_INT_MAX):
            continue
        found.append(raw)
    return found


def sheet_refs(text: str) -> list[str]:
    """Sheet names a formula qualifies references with, string literals masked out first."""
    masked = STRING_LITERAL.sub(lambda m: MASK * len(m.group()), text)
    return [(quoted or bare).replace("''", "'") for quoted, bare in SHEET_REF_RE.findall(masked)]


def external_reference(text: str) -> str | None:
    """The first operand naming another workbook, as `[Book.xlsx]Sheet!A1`, or None.

    Same predicate as recalc.py's gate, so the two scripts agree on what a link is: read
    the operands through the tokenizer, because a `[` inside a string literal is prose and
    a `Table[Column]` operand names no book at all.
    """
    try:
        tokens = Tokenizer(text if text.startswith("=") else "=" + text).items
    except Exception:
        return None  # an unparseable formula is recalc's to report, not this check's
    for token in tokens:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE or "!" not in token.value:
            continue
        if "[" in token.value.rsplit("!", 1)[0]:
            return token.value
    return None


def defined_names(wb) -> list[tuple[str | None, str, str]]:
    """(scope, name, expression) for every defined name, scope None for a workbook name.

    openpyxl files workbook-scoped names on the workbook and sheet-scoped ones on the
    sheet, so both tables have to be read. The scope stays part of the answer because the
    two can hold the same name pointing at different cells, which is a shape Excel allows
    and resolves by scope rather than by a single global table.
    """
    tables: list[tuple[str | None, object]] = [(None, wb.defined_names)]
    tables += [(ws.title, getattr(ws, "defined_names", {})) for ws in wb.worksheets]
    return [(scope, name, str(dn.value or "")) for scope, table in tables for name, dn in table.items()]


def external_names(wb) -> list[tuple[str, str, str]]:
    """(scope, name, linked reference) for every defined name pointing into another workbook.

    A cell reading `=ExternalRate` carries no `[` of its own, so the link sits one level
    down in the name and the per-cell scan walks past it.
    """
    found: list[tuple[str, str, str]] = []
    for scope, name, expression in defined_names(wb):
        linked = external_reference(expression)
        if linked:
            found.append((scope or "workbook", name, linked))
    return found


def area_cells(sheet: str, coord: str, expand_ranges: bool) -> frozenset[tuple[str, int, int]]:
    """The cells one reference operand covers, `$` already stripped, empty when it covers none.

    A range is skipped unless the caller asked for it expanded, and a range wider than
    RANGE_EXPAND_MAX is skipped either way, so a whole-column reference does not
    materialise a million cells to trace one dependency through.
    """
    try:
        if ":" not in coord:
            letter, line = coordinate_from_string(coord)
            return frozenset({(sheet, line, column_index_from_string(letter))})
        if not expand_ranges:
            return frozenset()
        first, last = coord.split(":", 1)
        c1, r1 = coordinate_from_string(first)
        c2, r2 = coordinate_from_string(last)
        i1, i2 = sorted((column_index_from_string(c1), column_index_from_string(c2)))
        r1, r2 = sorted((r1, r2))
        if (i2 - i1 + 1) * (r2 - r1 + 1) > RANGE_EXPAND_MAX:
            return frozenset()
        return frozenset((sheet, r, i) for r in range(r1, r2 + 1) for i in range(i1, i2 + 1))
    except Exception:
        return frozenset()  # a defined name or a whole-column range has no coordinates


def name_target(expression: str, expand_ranges: bool = False, scope: str | None = None, seen: frozenset[tuple[str | None, str]] = frozenset()) -> frozenset[tuple[str, int, int]]:
    """Every cell a defined name reads, as (sheet, row, column), empty when it reads none.

    A name is an alias for one or more areas (`Inputs!$B$2,Inputs!$B$4`) or a small formula
    over them (`Residual = ABS(Inputs!$B$2-Inputs!$B$3)`), and a cell reading the name
    depends on the same cells either way, so the whole body is read operand by operand
    rather than credited only when it is one bare reference: a `Checks` verdict written
    `=IF(Residual<0.01,"OK","FAIL")` does reach the solved input the residual subtracts.
    `scope` is the sheet a sheet-scoped name belongs to, which is both the home a name
    nested inside the body resolves against and the sheet an operand carrying no `!`
    sits on.
    """
    try:
        items = Tokenizer("=" + expression).items
    except Exception:
        return frozenset()
    cells: set[tuple[str, int, int]] = set()
    for t in items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        owner, bang, coord = t.value.rpartition("!")
        owner = owner.strip("'").replace("''", "'")
        nested = resolve_name(owner if bang else None, coord, scope, expand_ranges, seen)
        if nested is not None:
            cells |= nested
        elif bang:
            cells |= area_cells(owner.lower(), coord.replace("$", ""), expand_ranges)
        elif scope:
            cells |= area_cells(scope.lower(), coord.replace("$", ""), expand_ranges)
        # A bare coordinate in a workbook name reads on whichever sheet used the name, so
        # the same name means two cells on two sheets and the definition alone names none.
    return frozenset(cells)


def bare_link(text: str) -> bool:
    """A formula whose whole body is one cell reference, which is a link rather than a total."""
    return bool(LINK_FORMULA.match(text) or re.fullmatch(r"=\s*\$?[A-Za-z]{1,3}\$?\d+\s*", text))


def total_formula(text: str) -> bool:
    """True when a total row's formula is a SUM, an explicit addition or subtraction, or a bare link.

    Tokenising is what separates `=SUM(B3:B7)` from `=AVERAGE(B3:B7)`: both carry a
    parenthesis, so a substring test for one waves through every function ever written.
    """
    if bare_link(text):
        return True
    try:
        items = [t for t in Tokenizer(text).items if t.type != Token.WSPACE]
    except Exception:
        return False
    explicit = False  # a SUM or a +/- operator, so =(B3) or =1000 is not read as a total
    for t in items:
        if t.type == Token.FUNC:
            if t.subtype == Token.OPEN:
                if t.value.rstrip("(").strip().upper().removeprefix("_XLFN.") != "SUM":
                    return False
                explicit = True
        elif t.type in (Token.OP_PRE, Token.OP_IN):
            if t.value not in ("+", "-"):
                return False
            explicit = True
        elif t.type not in (Token.OPERAND, Token.PAREN, Token.SEP):
            return False
    return explicit


def formula_shape(text: str, row: int = 0, col: int = 0) -> str:
    """A formula reduced to its shape, so two cells of the same line item compare equal.

    References are rewritten relative to the cell that holds them (R1C1 terms), so a
    formula copied across a row keeps one shape while a copied formula whose reference
    was then edited by hand (`=D2*(1+$B$99)` beside `=C2*(1+$B$1)`) does not.
    """
    body = STRING_LITERAL.sub('""', text.upper())

    def relative(m: re.Match) -> str:
        ref = m.group(0)
        if ref.endswith("!"):
            return ref
        cabs, letters, rabs, digits = CELL_TOKEN.match(ref).groups()
        r, c = int(digits), column_index_from_string(letters)
        return (f"R{r}" if rabs else f"R[{r - row}]") + (f"C{c}" if cabs else f"C[{c - col}]")

    return re.sub(r"\s+", "", FAMILY_TOKEN.sub(relative, body))


_REFS: dict[tuple, frozenset] = {}
# Defined names of the workbook under audit as {(scope, name): expression}, lower-cased,
# scope None for a workbook name; set by audit().
_NAMES: dict[tuple[str | None, str], str] = {}


def resolve_name(owner: str | None, token: str, home: str | None, expand_ranges: bool = False, seen: frozenset[tuple[str | None, str]] = frozenset()) -> frozenset[tuple[str, int, int]] | None:
    """The cells a defined-name operand stands for, resolved in Excel's scope order.

    `Sheet!Name` is that sheet's own definition. A bare name is the formula's own sheet
    first and the workbook only if that sheet does not define it, which is the rule that
    keeps two sheets' `Rate` on their own cells instead of collapsing them into one; a
    name read from inside another name's body takes `home` from that name's scope, which
    is no sheet at all for a workbook name. None is reserved for a token no name defines,
    which is the answer that lets the caller read it as a coordinate; a name that stands
    for no cell resolves to an empty set.
    """
    scopes = (owner.lower(),) if owner is not None else ((home.lower(), None) if home else (None,))
    for scope in scopes:
        key = (scope, token.lower())
        expression = _NAMES.get(key)
        if expression is None:
            continue
        # A name that reaches itself, directly or through another name, stops here.
        return frozenset() if key in seen else name_target(expression, expand_ranges, scope, seen | {key})
    return None


def cell_refs(text: str, home: str, expand_ranges: bool = False) -> frozenset[tuple[str, int, int]]:
    """Every cell a formula reads, as (sheet, row, column), anchors ignored.

    Tokenising rather than pattern-matching is what keeps `LOG10(` and a quoted sheet name
    out of the answer; a reference with no sheet qualifier belongs to `home`. A range is
    skipped by default, because no axis header is one; the dependency walk asks for it
    expanded, since a residual that reads a solved input through `SUM` or `NPV` is the
    normal case, and a whole-column range is skipped rather than expanded. A defined name
    resolves under the same rule, area by area.
    """
    key = (text, home.lower(), expand_ranges)
    hit = _REFS.get(key)
    if hit is not None:
        return hit
    found: set[tuple[str, int, int]] = set()
    try:
        items = Tokenizer(text).items
    except Exception:
        items = []
    for t in items:
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        owner, bang, coord = t.value.rpartition("!")
        owner = owner.strip("'").replace("''", "'")
        sheet = (owner or home).lower()
        named = resolve_name(owner if bang else None, coord, home, expand_ranges)
        found |= named if named is not None else area_cells(sheet, coord.replace("$", ""), expand_ranges)
    hit = frozenset(found)
    _REFS[key] = hit
    return hit


def row_bands(group: list) -> list[list]:
    """Split cells of one formula shape into runs of adjacent rows, so stacked grids separate."""
    bands, band = [], []
    for c in sorted(group, key=lambda c: (c.row, c.column)):
        if band and c.row - band[-1].row > 1:
            bands.append(band)
            band = []
        band.append(c)
    return bands + [band] if band else bands


def detect_grids(ws, formula_cells: list, vsheet) -> list[dict]:
    """Sensitivity grids on one sheet, as {sheet, rows, cols, centre}.

    Shape is the whole test, so a grid is found wherever it sits: a filled rectangle of one
    formula shape, numeric headers down the column to its left and across the row above it,
    and every cell reading the two headers that cross on it. That last condition is what
    separates a two-way table from a schedule, which reads a period header and a text label,
    and it holds whether the axis values are typed on the sheet or linked from an inputs tab.
    """
    def numeric_at(row: int, col: int) -> bool:
        return is_number(ws.cell(row=row, column=col).value) or is_number(vsheet.cell(row=row, column=col).value)

    by_shape: dict[str, list] = {}
    for c in formula_cells:
        by_shape.setdefault(formula_shape(str(c.value), c.row, c.column), []).append(c)

    home = ws.title.lower()
    grids: list[dict] = []
    for group in by_shape.values():
        if len(group) < GRID_MIN_SIDE**2:
            continue
        for band in row_bands(group):
            rows = sorted({c.row for c in band})
            cols = sorted({c.column for c in band})
            if len(rows) < GRID_MIN_SIDE or len(cols) < GRID_MIN_SIDE or rows[0] == 1 or cols[0] == 1:
                continue
            if cols != list(range(cols[0], cols[-1] + 1)) or len(band) != len(rows) * len(cols):
                continue  # not a filled rectangle
            head_row, head_col = rows[0] - 1, cols[0] - 1
            if not all(numeric_at(r, head_col) for r in rows) or not all(numeric_at(head_row, c) for c in cols):
                continue
            if not all(
                {(home, c.row, head_col), (home, head_row, c.column)} <= cell_refs(str(c.value), ws.title)
                for c in band
            ):
                continue  # the cells do not vary along the two axes, so it is not a grid
            centre = None
            if len(rows) % 2 and len(cols) % 2:
                centre = ws.cell(row=rows[len(rows) // 2], column=cols[len(cols) // 2]).coordinate
            grids.append({"sheet": ws.title, "rows": rows, "cols": cols, "centre": centre})
    return grids


def subtracts(text: str, home: str, pair: set[tuple[str, int, int]]) -> bool:
    """True when a formula subtracts two given cells from each other, the shape a tie-out is written in.

    Reading the operands either side of the minus is what separates the tie-out from a
    formula that merely mentions both cells, and it holds whether the difference stands
    alone in column B or sits inside the verdict's own ABS.
    """
    try:
        items = [t for t in Tokenizer(text).items if t.type != Token.WSPACE]
    except Exception:
        return False
    for i in range(1, len(items) - 1):
        if items[i].type != Token.OP_IN or items[i].value != "-":
            continue
        ends: set[tuple[str, int, int]] = set()
        for side in (items[i - 1], items[i + 1]):
            if side.type == Token.OPERAND and side.subtype == Token.RANGE:
                ends |= cell_refs("=" + side.value, home)
        if ends == pair:
            return True
    return False


def designated_outputs(wb, values, grid: dict, cells: set[tuple[str, int, int]]) -> tuple[list[tuple], str]:
    """The output cells a workbook designates for one grid's centre, and how it named them.

    A grid centre is the base case, so it has to reproduce the model's own headline. Which
    cell that is has to come from the workbook rather than from whichever formula happens to
    mention the centre: a Checks row that carries a verdict and pairs the centre with one
    other cell is the designation the contract asks for, the same difference written
    elsewhere still names the output, and a data table names it in its own formula cell.
    Nothing else designates. A workbook names its WACC and its tax rate too, and a grid
    read against whichever named cell falls nearest the centre fails a correctly wired
    table and passes a disconnected one by coincidence; a grid that designates no output
    is reported unverified instead.
    """
    home, centre = grid["sheet"].lower(), grid["centre"]
    col, row = coordinate_from_string(centre)
    centre_ref = (home, row, column_index_from_string(col))

    def cached(sheet: str, coord: str):
        return values[sheet][coord].value if sheet in values.sheetnames else None

    paired: list[tuple] = []
    tied: list[tuple] = []
    for ws in wb.worksheets:
        on_checks = bool(CHECKS_SHEET.match(ws.title))
        for line in ws.iter_rows():
            verdict = on_checks and any(x.data_type == "f" and VERDICT_TOKEN.search(str(x.value)) for x in line)
            for c in line:
                if c.data_type != "f" or (ws.title.lower(), c.row, c.column) in cells:
                    continue
                refs = cell_refs(str(c.value), ws.title)
                if centre_ref not in refs:
                    continue
                others = [r for r in refs - {centre_ref} if r not in cells]
                if len(others) != 1:
                    continue
                sheet = next((s for s in wb.sheetnames if s.lower() == others[0][0]), None)
                if sheet is None:
                    continue
                coord = f"{get_column_letter(others[0][2])}{others[0][1]}"
                # A plain `=IF(<centre>><spot>,...)` on a model sheet pairs the centre with a
                # cell the grid was never meant to reproduce, so a designation is either a
                # Checks verdict row or the difference the tie-out is written as.
                if verdict:
                    paired.append((sheet, coord, cached(sheet, coord)))
                elif subtracts(str(c.value), ws.title, {centre_ref, others[0]}):
                    tied.append((sheet, coord, cached(sheet, coord)))
    if paired:
        return paired, "the check row pairing the centre"
    if tied:
        return tied, "a difference against the centre"

    # A data table writes its output into the cell where its two axis headers meet, which is
    # the one place a grid says for itself which number it varies. The cell it reads counts
    # whether it is addressed directly or through a defined name, since `=FairValue` there
    # designates as plainly as `=Model!B10`; a name nothing in the grid points at does not.
    home_sheet = next((s for s in wb.sheetnames if s.lower() == home), None)
    if home_sheet is None:
        return [], ""
    corner = wb[home_sheet].cell(row=grid["rows"][0] - 1, column=grid["cols"][0] - 1)
    if corner.data_type != "f":
        return [], ""
    axes = {(home, r, corner.column) for r in grid["rows"]} | {(home, corner.row, c) for c in grid["cols"]}
    outside = [r for r in cell_refs(str(corner.value), home_sheet) if r not in cells and r not in axes]
    if len(outside) != 1:
        return [], ""
    sheet = next((s for s in wb.sheetnames if s.lower() == outside[0][0]), None)
    if sheet is None:
        return [], ""
    coord = f"{get_column_letter(outside[0][2])}{outside[0][1]}"
    return [(sheet, coord, cached(sheet, coord))], f"the table formula cell {corner.coordinate}"


class Finding:
    def __init__(self, check: str, level: str, message: str):
        self.check, self.level, self.message = check, level, message
        self.cells: list[str] = []
        self.count = 0

    def add(self, sheet: str, coord: str, note: str = "") -> None:
        self.count += 1
        if len(self.cells) < EXAMPLES:
            self.cells.append(f"{sheet}!{coord}" + (f" ({note})" if note else ""))

    def as_dict(self) -> dict:
        return {"check": self.check, "level": self.level, "count": self.count, "message": self.message, "examples": self.cells}


class SheetNotFound(Exception):
    """A --sheet name the workbook does not have; auditing nothing is not a pass."""


def audit(path: Path, only_sheet: str | None, strict: bool = False) -> dict:
    wb = load_workbook(path, data_only=False)
    if only_sheet is not None and only_sheet not in wb.sheetnames:
        raise SheetNotFound(f"--sheet {only_sheet!r} is not in this workbook; it has {wb.sheetnames}")
    values = load_workbook(path, data_only=True)
    _REFS.clear()
    _NAMES.clear()
    for scope, name, expression in defined_names(wb):
        _NAMES[(scope.lower() if scope else None, name.lower())] = expression
    sheet_names_lower = {n.lower() for n in wb.sheetnames}
    referenced: set[str] = set()
    total_formulas = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.data_type == "f":
                    total_formulas += 1
                    referenced.update(ref.lower() for ref in sheet_refs(str(c.value)))
    findings: dict[str, Finding] = {}

    def f(check: str, level: str, message: str) -> Finding:
        if check not in findings:
            findings[check] = Finding(check, level, message)
        return findings[check]

    stats = {"formula_cells": 0, "numeric_constants": 0, "derived_constants": 0, "inputs_without_comment": 0,
             "formula_literals": 0, "formulas_wrong_color": 0, "formulas_general_format": 0,
             "formula_hardcodes": 0, "formula_family_outliers": 0, "hidden_sheets": 0,
             "formula_density": {"total_formulas": total_formulas, "per_sheet": {}}}
    density = stats["formula_density"]["per_sheet"]

    def source_comment(ws, c, note: str) -> None:
        """Blue declares an input wherever the cell sits, so the rule follows the colour."""
        if rgb(c.font) in BLUE and not note.strip():
            stats["inputs_without_comment"] += 1
            f("provenance", "warn", "blue input without a cell comment naming its source (an empty comment does not count)").add(ws.title, c.coordinate)

    grids: list[dict] = []
    solved: list[tuple[str, str, int, int]] = []

    if getattr(getattr(wb, "calculation", None), "iterate", False):
        f("iterative_calc", "warn", "iterative calculation is on, which usually means a circular reference; state which loop is intended in a note, or clear wb.calculation.iterate").add("workbook", "calcPr")

    for scope, name, linked in external_names(wb):
        f("references_external", "fail", "reference to another workbook; recalc cannot resolve it").add(scope, name, f"defined name, {linked}")

    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        if ws.sheet_state != "visible":
            stats["hidden_sheets"] += 1
            hidden = f("hidden_sheet", "warn", "sheet is hidden; a reader cannot check what it holds, and veryHidden cannot be unhidden from the Excel UI at all")
            if ws.sheet_state == "veryHidden":
                hidden.level = "fail"
            hidden.add(ws.title, "A1", ws.sheet_state)
        cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
        formula_cells = [c for c in cells if c.data_type == "f"]
        if not formula_cells and ws.title.lower() not in referenced:
            # `referenced` reads Sheet!Cell text only, so a sheet consumed through a workbook
            # defined name lands here too: the density score and the colour verdict stay off,
            # but a blue cell is a declared input on any sheet and still owes its source.
            for c in cells:
                if not is_number(c.value):
                    continue
                note = (c.comment.text or "") if c.comment else ""
                if SOLVED_COMMENT.match(note):  # a Solved: value parked on a sheet nothing reads is still an orphan
                    solved.append((ws.title, c.coordinate, c.row, c.column))
                source_comment(ws, c, note)
            continue  # raw data sheet, so no density score and no colour verdict on its numbers
        # A formula-free sheet the model reads is raw data left as values or a sheet of pure
        # inputs, and the audit cannot tell which, so it is not a calculation sheet: no
        # density score, no colour verdict on its numbers. A blue cell on it is still an
        # input and still owes its comment.
        calculation_sheet = bool(formula_cells)
        formula_rows = {c.row for c in formula_cells}
        formula_cols = {c.column for c in formula_cells}
        stats["formula_cells"] += len(formula_cells)

        # Density: of the numbers a calculation sheet shows, how many it works out. A blue
        # cell is a declared input and is not counted against the sheet, and neither is a
        # period header, which is what the sheet is ruled with rather than something it
        # computes; every other typed number is one the model was supposed to compute.
        vsheet = values[ws.title]
        header_rows = period_header_rows(cells, formula_rows)
        computed = [c for c in formula_cells if evaluates_to_number(str(c.value), vsheet[c.coordinate].value)]
        typed = [c for c in cells if c.data_type != "f" and is_number(c.value)
                 and rgb(c.font) not in BLUE and c.row not in header_rows]
        scored = len(computed) + len(typed)
        share = round(len(computed) / scored, 3) if scored else None
        density[ws.title] = {"formulas": len(formula_cells), "numeric_cells": scored, "share": share}
        if calculation_sheet and scored >= DENSITY_MIN_CELLS and share < DENSITY_WARN:
            level = "fail" if share < DENSITY_FAIL else "warn"
            reported = f("formula_density", level, "calculation sheet shows more typed numbers than it works out; every derived value is a formula, and only a blue declared input is exempt")
            reported.level = "fail" if level == "fail" else reported.level
            reported.add(ws.title, "A1", f"{len(computed)} of {scored} numeric cells are formulas")
        grids.extend(detect_grids(ws, formula_cells, vsheet))

        for c in cells:
            value = c.value
            if c.data_type == "f":
                text = str(value)
                cached = vsheet[c.coordinate].value
                if isinstance(cached, str) and ERROR_VALUE.match(cached):
                    f("cached_error", "fail", "formula evaluates to an error; a Checks verdict built on it reads as a verdict while meaning nothing, and #ERROR! usually means a construct the calculation engine cannot evaluate (a range qualified with its sheet at both ends, --(array) coercion)").add(ws.title, c.coordinate, f"{cached} {text[:60]}")
                for prefix, name in FUNC_RE.findall(text):
                    name = name.upper()
                    if name in DYNAMIC_ARRAY:
                        f("functions", "fail", "dynamic-array or 365-only function; older Excel and LibreOffice show #NAME?").add(ws.title, c.coordinate, name)
                    elif name in NEEDS_XLFN and not prefix:
                        f("functions_prefix", "fail", "function needs the _xlfn. prefix when written by openpyxl or Excel shows #NAME?").add(ws.title, c.coordinate, name)
                    elif name in NETWORK:
                        f("functions_network", "fail", "network or external-data function; never emit these").add(ws.title, c.coordinate, name)
                    elif name in VOLATILE:
                        f("functions_volatile", "warn", "volatile function; replace with direct references unless required").add(ws.title, c.coordinate, name)
                for ref in sheet_refs(text):
                    if ref and ref.lower() not in sheet_names_lower and not ref.startswith("["):
                        f("references", "fail", "cross-sheet reference to a sheet that does not exist").add(ws.title, c.coordinate, ref)
                linked = external_reference(text)
                if linked:
                    f("references_external", "fail", "reference to another workbook; recalc cannot resolve it").add(ws.title, c.coordinate, linked)
                if LITERAL_FORMULA.match(text):
                    stats["formula_literals"] += 1
                    f("formula_literal", "fail", "formula is a bare constant; either type the number as an input or write the real formula").add(ws.title, c.coordinate, text)
                else:
                    # Every formula, wherever it sits: the header and label exclusions below
                    # are about numbers somebody typed, and `A2 = B2*0.42` in a label column
                    # hides the same assumption it would hide anywhere else.
                    literals = hardcoded_literals(text, ws.title)
                    if literals:
                        stats["formula_hardcodes"] += 1
                        f("formula_hardcode", "fail", "number typed inside a formula; put it in a labelled input cell and reference it, or the assumption is hidden where nobody will find it").add(ws.title, c.coordinate, ", ".join(literals[:4]))
                color = rgb(c.font)
                if color in BLUE:
                    f("colors_formula_blue", "warn", "formula cell styled as an input (blue); use black, or green for cross-sheet links").add(ws.title, c.coordinate)
                elif color not in BLACK and color not in GREEN and not (color in RED and "[" in text) and evaluates_to_number(text, cached):
                    # a text formula (a header built with &, a label) wears the header's colour
                    stats["formulas_wrong_color"] += 1
                    f("formula_font_color", "fail", "formula font off the colour code; black for a formula, green for a cross-sheet link, red only for a link to another workbook").add(ws.title, c.coordinate, str(color))
                if c.number_format == "General" and c.row != 1 and evaluates_to_number(text, cached):
                    stats["formulas_general_format"] += 1
                    f("formula_number_format", "fail" if strict else "warn", "numeric formula left on General; give it the format its line item reads in (currency, percent or multiple)").add(ws.title, c.coordinate)
                if c.comment and SOLVED_COMMENT.match(c.comment.text or ""):
                    f("solved_formula", "warn", "a Solved: comment on a formula; a solved value is the typed result of root-finding, and a number a formula produces stays a formula with no Solved: claim").add(ws.title, c.coordinate, text[:60])
            elif is_number(value):
                stats["numeric_constants"] += 1
                note = (c.comment.text or "") if c.comment else ""
                is_solved = bool(SOLVED_COMMENT.match(note))
                in_formula_line = c.row in formula_rows or c.column in formula_cols
                header_like = c.row == 1 or (isinstance(ws.cell(row=c.row, column=1).value, str) and TOTAL_LABEL.search(str(ws.cell(row=c.row, column=1).value) or ""))
                if in_formula_line and not header_like and not is_solved:
                    # a constant sitting on a line that is otherwise calculated
                    row_formulas = [p for p in formula_cells if p.row == c.row and p.column > 1]
                    if row_formulas and any(p.column > c.column for p in row_formulas) and any(p.column < c.column for p in row_formulas):
                        stats["derived_constants"] += 1
                        f("formulas", "fail", "hardcoded number between formulas on a calculated line; write the formula instead").add(ws.title, c.coordinate, str(value))
                color = rgb(c.font)
                if color not in BLUE and c.row != 1 and not header_like and calculation_sheet:
                    f("colors_input", "warn", "numeric input not styled blue (font color 0000FF)").add(ws.title, c.coordinate)
                source_comment(ws, c, note)
                if is_solved:
                    solved.append((ws.title, c.coordinate, c.row, c.column))
                    if color not in BLUE:
                        f("solved_color", "warn", "a Solved: value is a typed input and is styled blue; any other colour hides that the reader may nudge it").add(ws.title, c.coordinate)
                    if len(SOLVED_TARGET.findall(note)) < 2 or not SOLVED_RESOLVE.search(note):
                        f("solved_comment", "fail" if strict else "warn", "a Solved: comment reads `Solved: <what> at which <target cell> equals <spot cell>; re-solve with <command>`; one that names fewer than two cells or no command cannot be re-solved or checked").add(ws.title, c.coordinate, note.strip()[:60])
                if c.number_format == "General" and isinstance(value, float) and abs(value) < 1:
                    f("formats_percent", "warn", "fraction with General format; use 0.0% if it is a rate").add(ws.title, c.coordinate, str(value))

        # totals: a blue cell is an input whatever its label says ("Total debt" off a balance sheet)
        for c in cells:
            if c.column == 1 and isinstance(c.value, str) and TOTAL_LABEL.search(c.value):
                for p in [p for p in cells if p.row == c.row and p.column > 1]:
                    if is_number(p.value) and rgb(p.font) not in BLUE:
                        f("totals", "fail", "total row holds a typed number; use =SUM(...) over the block above, or style it blue if it is an input").add(ws.title, p.coordinate)
                    elif p.data_type == "f" and not total_formula(str(p.value)):
                        f("totals_formula", "warn", "total row formula is neither a SUM, an explicit addition or subtraction, nor a bare link to the cell that holds the figure").add(ws.title, p.coordinate)

        # frozen rows: a formula copied across periods shifts its relative references, so
        # identical text in three or more adjacent period cells means every period reads
        # the same cell. recalc cannot see this; it is the model that is wrong, not a formula.
        by_row: dict[int, list] = {}
        for c in formula_cells:
            by_row.setdefault(c.row, []).append(c)
        for row_cells in by_row.values():
            row_cells.sort(key=lambda c: c.column)
            run: list = []
            for c in row_cells + [None]:
                if c is not None and run and c.column == run[-1].column + 1 and str(c.value) == str(run[-1].value):
                    run.append(c)
                    continue
                if len(run) >= 3 and RELATIVE_REF.search(str(run[0].value)):
                    f("frozen_row", "warn", "identical formula text across adjacent period columns; a copied formula shifts its relative references, so these all read one period").add(ws.title, f"{run[0].coordinate}:{run[-1].coordinate}", str(run[0].value)[:60])
                run = [c] if c is not None else []

        # formula families: the inverse of frozen_row. One line item is one formula copied
        # across, so a row that agrees on a shape everywhere but one cell has been edited in
        # place. A run of one shape at either end is exempt, because it differs by design:
        # the historical periods link to a data sheet, and the last column carries an exit
        # or a terminal value. Only the interior is copied, so only the interior is diagnosable.
        for row_cells in by_row.values():
            family = [c for c in row_cells if c.column > 1]
            if len(family) < FAMILY_MIN_CELLS:
                continue
            shapes = {c.coordinate: formula_shape(str(c.value), c.row, c.column) for c in family}
            dominant, hits = Counter(shapes.values()).most_common(1)[0]
            outliers = [c for c in family if shapes[c.coordinate] != dominant]
            if hits / len(family) < FAMILY_DOMINANT or not 1 <= len(outliers) <= FAMILY_MAX_OUTLIERS:
                continue
            ordered = sorted(family, key=lambda c: c.column)
            edges: set[int] = set()
            for end in (ordered, ordered[::-1]):
                first = shapes[end[0].coordinate]
                if first == dominant:
                    continue
                for c in end:
                    if shapes[c.coordinate] != first:
                        break
                    edges.add(c.column)
            for c in outliers:
                if c.column in edges:
                    continue
                stats["formula_family_outliers"] += 1
                f("formula_family", "warn", "formula does not match the shape the rest of its row shares; a line item is one formula copied across periods, so check whether this cell was edited by hand").add(ws.title, c.coordinate, dominant[:60])

        # layout
        used_cols = sorted({c.column for c in cells})
        # Reading column_dimensions[letter] creates a default entry, so ask for the keys first.
        sized = {k for k, d in ws.column_dimensions.items() if d.width and d.customWidth}
        default_widths = [get_column_letter(col) for col in used_cols if get_column_letter(col) not in sized]
        if len(default_widths) > max(1, len(used_cols) // 2):
            f("layout_widths", "warn", "most used columns keep the default width; set column_dimensions[...].width").add(ws.title, ",".join(default_widths[:6]))
        if ws.freeze_panes is None and ws.max_row > 25:
            f("layout_freeze", "warn", "no frozen header row on a long sheet; set ws.freeze_panes").add(ws.title, "A1")
        if ws.sheet_view.showGridLines:
            f("layout_gridlines", "info", "gridlines visible; model sheets read cleaner with ws.sheet_view.showGridLines = False").add(ws.title, "A1")

    visible = [ws for ws in wb.worksheets if ws.sheet_state == "visible"]
    if len(wb.worksheets) >= 3 and visible:
        cover = visible[0]
        if not any(c.data_type == "f" and LINK_FORMULA.match(str(c.value)) for row in cover.iter_rows() for c in row):
            f("cover_static", "warn", "first visible sheet links to no output cell, so it is a title page or the model itself rather than a cover; the sheet a reader meets first carries tiles that link to the output cells they report, or they go stale the moment an input changes").add(cover.title, "A1")

    for ws in wb.worksheets:
        if not CHECKS_SHEET.match(ws.title):
            continue
        texts = [str(c.value).upper() for row in ws.iter_rows() for c in row if c.data_type == "f"]
        if not any(VERDICT_TOKEN.search(x) for x in texts):
            f("checks_rollup", "fail" if strict else "warn", "Checks sheet carries no verdict formula; every row tests with a live =IF(...,\"OK\",\"FAIL\") or WARN formula in column C").add(ws.title, "C5")
        elif not any("COUNTIF" in x and "FAIL" in x for x in texts):
            f("checks_rollup", "fail" if strict else "warn", "Checks sheet has no roll-up counting FAIL verdicts; the Overall row is what the delivery reads back").add(ws.title, "A1")

    if grids:
        if not any(CHECKS_SHEET.match(n) for n in wb.sheetnames):
            f("checks_missing", "warn", "workbook has a sensitivity grid and no Checks sheet; the direction of each axis belongs in a live check row, so a grid that moves the wrong way is caught by the workbook itself").add(grids[0]["sheet"], grids[0]["centre"] or "A1")
        in_grid = {(g["sheet"].lower(), r, c) for g in grids for r in g["rows"] for c in g["cols"]}

        def unverified(g: dict, coord: str, why: str) -> None:
            f("sensitivity_unverified", "warn", "sensitivity grid centre could not be checked against the output it varies; an unchecked grid reads as a valid table whether or not it is wired to anything").add(g["sheet"], coord, why)

        for g in grids:
            if not g["centre"]:
                corner = f"{get_column_letter(g['cols'][0])}{g['rows'][0]}"
                unverified(g, corner, "grid has an even side, so no cell holds the base case; use odd dimensions")
                continue
            centre = values[g["sheet"]][g["centre"]].value
            if not is_number(centre):
                unverified(g, g["centre"], "centre holds no cached value; run recalc.py before audit.py")
                continue
            outputs, source = designated_outputs(wb, values, g, in_grid)
            if not outputs:
                unverified(g, g["centre"], "no check row, tie-out or output designation found; add a Checks row comparing the centre with the output the grid varies")
                continue
            numeric = [o for o in outputs if is_number(o[2])]
            if not numeric:
                unverified(g, g["centre"], f"{source} holds no cached value; run recalc.py before audit.py")
                continue

            def off_by(other: float, centre: float = centre) -> float:
                return abs(other - centre) / (abs(centre) or abs(other) or 1.0)

            # A check row is a designation and every one has to hold; a workbook can write
            # more than one difference against the centre, so the closest of those is taken.
            # A table formula cell names one output either way.
            judged = max(numeric, key=lambda o: off_by(o[2])) if source.startswith("the check row") else min(numeric, key=lambda o: off_by(o[2]))
            gap = off_by(judged[2])
            if gap > GRID_CENTRE_TOLERANCE:
                f("sensitivity_centre", "fail", "sensitivity grid centre does not reproduce the output it varies; the centre is the base case, so a grid centred off it or wired to a dead cell reads as a valid table and means nothing").add(
                    g["sheet"], g["centre"], f"{centre:.6g}, {source} names {judged[0]}!{judged[1]} {judged[2]:.6g}, off by {gap:.1%}")

    if solved:
        dependents: dict[tuple[str, int, int], set[tuple[str, int, int]]] = {}
        verdicts: set[tuple[str, int, int]] = set()
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.data_type != "f":
                        continue
                    text = str(c.value)
                    key = (ws.title.lower(), c.row, c.column)
                    for ref in cell_refs(text, ws.title, expand_ranges=True):
                        dependents.setdefault(ref, set()).add(key)
                    if CHECKS_SHEET.match(ws.title) and VERDICT_TOKEN.search(text):
                        verdicts.add(key)
        for sheet, coord, r, col in solved:
            start = (sheet.lower(), r, col)
            seen = {start}
            frontier = [start]
            reached = False
            while frontier and not reached:
                nxt = []
                for node in frontier:
                    for dep in dependents.get(node, ()):
                        if dep in seen:
                            continue
                        seen.add(dep)
                        if dep in verdicts:
                            reached = True
                            break
                        nxt.append(dep)
                    if reached:
                        break
                frontier = nxt
            if not reached:
                f("solved_unchecked", "fail", "a Solved: input that no Checks verdict depends on; add a residual row (forward value at the solved input minus its target, with a tolerance) so the workbook fails when the target moves and the value no longer fits").add(sheet, coord)

    result = [x.as_dict() for x in findings.values()]
    result.sort(key=lambda d: {"fail": 0, "warn": 1, "info": 2}[d["level"]])
    fails = sum(1 for d in result if d["level"] == "fail")
    return {"status": "fail" if fails else "pass", "file": str(path), "stats": stats, "findings": result}


USAGE = "usage: audit.py <file.xlsx> [--strict] [--sheet NAME]"
TAKES_VALUE = {"--strict": False, "--sheet": True}


def fail(message: str) -> None:
    print(json.dumps({"status": "error", "message": message}))
    sys.exit(1)


def parse_args(argv: list[str]) -> tuple[list[str], dict[str, str | bool]]:
    """Positional arguments and flags, refusing any token TAKES_VALUE does not name.

    A misspelt flag used to fall through unread, so `--stirct` audited without the
    gate it asked for and exited 0 on a workbook that fails.
    """
    args: list[str] = []
    flags: dict[str, str | bool] = {}
    i = 0
    while i < len(argv):
        token = argv[i]
        if not token.startswith("--"):
            args.append(token)
            i += 1
        elif token not in TAKES_VALUE:
            fail(f"unknown option: {token}; {USAGE}")
        elif token in flags:
            fail(f"repeated option: {token}; {USAGE}")
        elif not TAKES_VALUE[token]:
            flags[token] = True
            i += 1
        elif i + 1 >= len(argv) or argv[i + 1].startswith("--"):
            fail(f"{token} requires a value; {USAGE}")
        else:
            flags[token] = argv[i + 1]
            i += 2
    return args, flags


def main(argv: list[str]) -> None:
    if "-h" in argv or "--help" in argv:
        print(__doc__.strip())
        sys.exit(0)
    args, flags = parse_args(argv)
    if not args:
        fail(USAGE)
    path = Path(args[0]).expanduser().resolve()
    try:
        report = audit(path, flags.get("--sheet"), "--strict" in flags)
    except SheetNotFound as exc:
        print(json.dumps({"status": "error", "file": str(path), "message": str(exc)}))
        sys.exit(1)
    except Exception as exc:  # a workbook this script cannot open is a report, not a traceback
        print(json.dumps({"status": "error", "file": str(path), "message": f"{type(exc).__name__}: {exc}"}))
        sys.exit(1)
    print(json.dumps(report, indent=2))
    if "--strict" in flags and report["status"] == "fail":
        sys.exit(1)


if __name__ == "__main__":
    main(sys.argv[1:])
