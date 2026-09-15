#!/usr/bin/env python3
"""Profile a messy uploaded sheet before cleaning it: header row, column types, data quality.

Usage:
    python profile.py <file.xlsx|file.csv> [--sheet NAME] [--max-rows N] [--delimiter X]

Prints one JSON object. Per sheet: the detected header row, then for every column
`name`, `type`, `confidence`, `missing_pct`, `unique` and up to five `samples`, then
`issues` in the shape audit.py uses (`check`, `level`, `count`, `examples`).

Types:

    text number percent currency date boolean identifier mixed empty

A column of bare numbers carrying a percent or currency number format is reported as
`percent` or `currency`. The cells hold plain floats, so the format is the only thing
that says what those floats mean.

Issues:

    blank_header      a column the header row does not name
    duplicate_header  one name on two columns
    empty_rows        a row where every cell is blank
    empty_columns     a column where every cell is blank
    duplicate_rows    rows identical once case and spacing are ignored
    subtotal_rows     a total or subtotal label in the first cells of a row
    number_as_text    a number stored as a string, or one carrying a stray $ , % or ( )
    mixed_dates       more than one date format down one column, reading a real date cell
                      by the format it displays in and a string by the shape it was typed in
    whitespace        a leading, trailing or doubled space
    casing            one value spelled in several casings in a categorical column
    mojibake          UTF-8 read as Latin-1 (`Ã©`, `â€™`), a replacement character left
                      by a decode that lost the original byte, or a non-printing character
    error_cells       `#REF!` and the other spreadsheet errors

The header row is chosen by scoring the first 30 rows: distinct text cells score, a
filled row underneath scores, and numeric-looking or near-empty rows are penalised.

`--sheet` names one sheet of the workbook and is an error when the workbook has no
sheet by that name, rather than a success that profiled nothing.

`--max-rows` caps the scan at 5,000 rows a sheet and wants a whole number of at
least one row. A sheet cut short reports `truncated: true`, so `data_rows`, the column
types and every issue count on it describe the rows that were read and not the whole
file; raise the cap and re-run before proposing a clean-up of a longer upload.

A `.csv` is decoded as UTF-8 and, when that fails, as cp1252 then latin-1, and the one
that worked is reported as `encoding` beside `file`. A column that reads as nonsense is
then a file written in some other encoding, not a file this script mangled on the way in.

Its separator is sniffed from the same decoded sample, out of `,` `;` tab and `|`, and
reported as `delimiter`. A European export written with semicolons therefore profiles as
its real columns rather than as one text column holding every field of a row. `--delimiter`
names the separator outright, one character, for the file a sniff reads wrong.

This script never modifies the file. Read it, put the proposed fixes to the user, and
only then change anything.
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import Counter, defaultdict
from collections.abc import Iterable
from datetime import date, datetime
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

BLANK_TOKENS = {"", "-", "--", "n/a", "na", "null", "none", "nan"}
# The shape audit.py and recalc.py match, so the bare #ERROR! a calculation engine caches
# for a construct it cannot evaluate counts here too, as does a newer #BLOCKED!/#CALC!.
# The trailing ! or ? stays required, and the case stays significant, because this runs
# over every string in an upload where `#ref` is somebody's note rather than an error.
ERROR_CELL = re.compile(r"^#(?:N/A|[A-Z0-9/]+[!?])$")
TOTAL_LABEL = re.compile(r"\b(total|subtotal|sub-total|grand total|sum)\b", re.I)
ID_WORDS = {"id", "ids", "code", "sku", "zip", "postal", "phone", "cusip", "isin", "ticker", "account", "ref", "symbol"}
BOOLEANS = {"true", "false", "yes", "no", "y", "n"}
CURRENCY_SIGN = re.compile(r"[$€£¥₹]|\b(?:usd|eur|gbp|jpy|cny|rmb|hkd)\b", re.I)
LEADING_ZERO = re.compile(r"^0\d+$")
# Number formats. `[$-409]` is a locale id whose own $ means nothing, so it goes before the
# currency test; a quoted or escaped literal shows but does not scale, so `0.0" %"` goes
# before the percent test while `"$"#,##0.00` keeps its sign for the currency one.
FORMAT_LOCALE = re.compile(r"\[\$?-[^\]]*\]")
FORMAT_LITERAL = re.compile(r'"[^"]*"|\\.')
CURRENCY_SYMBOL = re.compile(r"[$€£¥₹]")
CSV_ENCODINGS = ("utf-8-sig", "cp1252", "latin-1")
# The separators an export actually writes. Left unrestricted, the sniffer settles on any
# punctuation that repeats evenly down the sample, so a column of ISO dates is enough to
# split a single-column file on `-`.
CSV_DELIMITERS = ",;\t|"
CSV_SNIFF_BYTES = 64 * 1024
# The two-byte openers a UTF-8 string picks up when it is decoded as Latin-1, plus a BOM,
# plus U+FFFD: a replacement character is a byte some decoder threw away, so the text it
# sits in already lost something, whether the loss happened here or upstream.
MOJIBAKE = re.compile("\u00c3[\u0080-\u00ff]|\u00c2[\u00a0-\u00bf]|\u00e2\u20ac|\ufeff|\ufffd")
NONPRINTING = re.compile("[\x00-\x08\x0b-\x1f\x7f\u200b-\u200f]")
DATE_SHAPES = [
    (re.compile(r"^\d{4}-\d{1,2}-\d{1,2}$"), "yyyy-mm-dd"),
    (re.compile(r"^\d{4}/\d{1,2}/\d{1,2}$"), "yyyy/mm/dd"),
    (re.compile(r"^\d{1,2}/\d{1,2}/\d{2}$"), "m/d/yy"),
    (re.compile(r"^\d{1,2}/\d{1,2}/\d{4}$"), "m/d/yyyy"),
    (re.compile(r"^\d{1,2}-\d{1,2}-\d{4}$"), "d-m-yyyy"),
    (re.compile(r"^\d{1,2}\.\d{1,2}\.\d{4}$"), "d.m.yyyy"),
    (re.compile(r"^\d{1,2}[- ][A-Za-z]{3,9}[- ]\d{2,4}$"), "d-mmm-yy"),
    (re.compile(r"^[A-Za-z]{3,9} \d{1,2},? \d{4}$"), "mmmm d, yyyy"),
]
# One sheet as it is read: its rows, whether the cap cut them, what each column's number
# format declares, and the number format of every date cell by column and sheet row.
SheetRead = tuple[list[list], bool, dict[int, str], dict[int, dict[int, str]]]
HEADER_SCAN = 30
SAMPLES = 5
EXAMPLES = 25
CATEGORICAL_MAX = 50


def norm(value) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        return str(value).strip()
    return re.sub(r"\s+", " ", value.replace("\u00a0", " ")).strip()


def is_blank(value) -> bool:
    return norm(value).lower() in BLANK_TOKENS


def parse_number(text: str) -> float | None:
    """A number as a human types it: currency signs, thousands commas, percent, `(12)` for -12."""
    s = text.strip()
    negative = s.startswith("(") and s.endswith(")")
    if negative:
        s = s[1:-1].strip()
    s = CURRENCY_SIGN.sub("", s).replace(",", "").replace(" ", "").lstrip("+")
    percent = s.endswith("%")
    if percent:
        s = s[:-1]
    try:
        value = float(s)
    except ValueError:
        return None
    if percent:
        value /= 100
    return -value if negative else value


def format_kind(number_format: str | None) -> str:
    """What a cell's number format says its bare numbers mean, or "" when it says nothing."""
    fmt = FORMAT_LOCALE.sub("", number_format or "")
    if "%" in FORMAT_LITERAL.sub("", fmt):
        return "percent"
    if CURRENCY_SYMBOL.search(fmt):
        return "currency"
    return ""


def date_shape(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return "datetime"
    text = norm(value)
    for pattern, name in DATE_SHAPES:
        if pattern.match(text):
            return name
    return None


def date_format(number_format: str | None) -> str:
    """A date cell's display format, normalised: a locale id and the case of the tokens
    say nothing about how the date reads, so `[$-409]MM/DD/YYYY` and `mm/dd/yyyy` are one
    format rather than two."""
    return FORMAT_LOCALE.sub("", number_format or "").strip().lower()


def date_written(value, number_format: str | None) -> str | None:
    """How one cell writes its date: a real date cell by the format it displays in, a
    string by the shape it was typed in, or None when it is not a date at all.

    openpyxl hands back a `datetime` for every real date cell whatever its number format,
    so a column shown half as `mm/dd/yyyy` and half as `yyyy-mm-dd` looked uniform and
    never reached `mixed_dates`. The two kinds keep separate tokens, because a date cell
    sitting beside a string that displays identically is the mix hardest to see by eye.
    """
    shape = date_shape(value)
    if shape != "datetime":
        return shape
    return f"datetime {date_format(number_format)}" if number_format else shape


def header_score(row: list, following: list, index: int) -> float:
    """How much a row reads like a header: named columns, data underneath, nothing computed."""
    filled = [norm(v) for v in row if not is_blank(v)]
    if len(filled) < 2:
        return -1e6
    texty = sum(1 for v in filled if re.search(r"[A-Za-z]", v)) / len(filled)
    numeric = sum(1 for v in filled if parse_number(v) is not None) / len(filled)
    unique = len(set(v.lower() for v in filled)) / len(filled)
    below = sum(1 for v in following if not is_blank(v)) / max(1, len(filled))
    score = 3 * texty + 2 * unique + 1.5 * min(below, 1.0) + 0.25 * len(filled) - 2 * numeric
    if any(TOTAL_LABEL.search(v) for v in filled):
        score -= 2
    if index and len(filled) == 2:
        score -= 1.5  # a title or a stray label/value pair, not a header
    return score


def detect_header(rows: list[list]) -> int:
    scored = [(header_score(row, rows[i + 1] if i + 1 < len(rows) else [], i), -i, i)
              for i, row in enumerate(rows[:HEADER_SCAN])]
    best = max(scored, default=(0, 0, 0))
    return best[2] if best[0] > -1e5 else 0


def classify(name: str, values: list, declared: str = "") -> tuple[str, float]:
    """Column type from the evidence its values carry, in order of how specific the signal is.

    `declared` is what the column's number format says (see `format_kind`), which for a
    column of bare floats is the only evidence there is.
    """
    if not values:
        return "empty", 1.0
    texts = [norm(v) for v in values]
    total = len(texts)
    words = set(re.split(r"[^a-z0-9]+", name.lower()))
    named_id = bool(words & ID_WORDS)
    zero_id = any(LEADING_ZERO.match(t) for t in texts)
    ratio = {
        "boolean": sum(1 for t in texts if t.lower() in BOOLEANS) / total,
        "number": sum(1 for t in texts if parse_number(t) is not None) / total,
        "percent": sum(1 for t in texts if t.endswith("%") or t.lower().endswith("bps")) / total,
        "currency": sum(1 for t in texts if CURRENCY_SIGN.search(t)) / total,
        "date": sum(1 for v in values if date_shape(v)) / total,
    }
    if ratio["boolean"] >= 0.9:
        return "boolean", round(ratio["boolean"], 2)
    # An identifying name, or a leading zero, says what a column is for rather than what it
    # holds, so both yield to the value evidence below that the column is a quantity:
    # `Account Balance` is currency, while `Account ID` is still an identifier.
    quantity = (
        (declared and ratio["number"] >= 0.8)
        or ratio["percent"] >= 0.6
        or (ratio["currency"] >= 0.4 and ratio["number"] >= 0.7)
        or ratio["date"] >= 0.8
        or any((n := parse_number(t)) is not None and isfinite(n) and n != int(n) for t in texts)
    )
    if (named_id or zero_id) and not quantity:
        return "identifier", 0.9 if named_id else 0.75
    # A declared format is the author's own statement of what the numbers are, so it
    # outranks the string shapes below, which can only see a % or a $ someone typed.
    if declared and ratio["number"] >= 0.8:
        return declared, 0.95
    if ratio["percent"] >= 0.6:
        return "percent", round(ratio["percent"], 2)
    if ratio["currency"] >= 0.4 and ratio["number"] >= 0.7:
        return "currency", round(min(0.95, (ratio["currency"] + ratio["number"]) / 2), 2)
    if ratio["date"] >= 0.8:
        return "date", round(ratio["date"], 2)
    if ratio["number"] >= 0.85:
        return "number", round(ratio["number"], 2)
    if max(ratio["number"], ratio["date"]) >= 0.3:
        return "mixed", round(max(ratio["number"], ratio["date"]), 2)
    return "text", 0.7


class Issue:
    def __init__(self, check: str, level: str):
        self.check, self.level = check, level
        self.count = 0
        self.examples: list[str] = []

    def add(self, where: str, note: str = "") -> None:
        self.count += 1
        if len(self.examples) < EXAMPLES:
            self.examples.append(where + (f" ({note})" if note else ""))

    def as_dict(self) -> dict:
        return {"check": self.check, "level": self.level, "count": self.count, "examples": self.examples}


def profile_sheet(name: str, rows: list[list], truncated: bool, formats: dict[int, str], date_formats: dict[int, dict[int, str]]) -> dict:
    issues: dict[str, Issue] = {}

    def issue(check: str, level: str) -> Issue:
        return issues.setdefault(check, Issue(check, level))

    header_index = detect_header(rows)
    header = rows[header_index] if rows else []
    width = max((len(r) for r in rows), default=0)
    data = rows[header_index + 1:]

    names: list[str] = []
    seen: dict[str, int] = {}
    for i in range(width):
        raw = norm(header[i]) if i < len(header) else ""
        coord = f"{get_column_letter(i + 1)}{header_index + 1}"
        if not raw:
            issue("blank_header", "warn").add(coord)
            raw = f"column_{get_column_letter(i + 1)}"
        elif raw.lower() in seen:
            issue("duplicate_header", "warn").add(coord, raw)
        seen[raw.lower()] = seen.get(raw.lower(), 0) + 1
        names.append(raw)
        cell = header[i] if i < len(header) else None
        if isinstance(cell, str):
            if cell != cell.strip() or "  " in cell:
                issue("whitespace", "info").add(coord, repr(cell)[:30])
            if MOJIBAKE.search(cell) or NONPRINTING.search(cell):
                issue("mojibake", "warn").add(coord, norm(cell)[:30])

    keys: dict[tuple, int] = {}
    for r, row in enumerate(data):
        sheet_row = header_index + 2 + r
        cells = [row[i] if i < len(row) else None for i in range(width)]
        if all(is_blank(v) for v in cells):
            issue("empty_rows", "info").add(f"A{sheet_row}")
            continue
        key = tuple(norm(v).lower() for v in cells)
        if key in keys:
            issue("duplicate_rows", "warn").add(f"A{sheet_row}", f"same as row {keys[key]}")
        else:
            keys[key] = sheet_row
        if any(TOTAL_LABEL.search(norm(v)) for v in cells[:3] if isinstance(v, str)):
            issue("subtotal_rows", "warn").add(f"A{sheet_row}", norm(cells[0])[:30])
        for i, value in enumerate(cells):
            coord = f"{get_column_letter(i + 1)}{sheet_row}"
            if not isinstance(value, str):
                continue
            if ERROR_CELL.match(value.strip()):
                issue("error_cells", "warn").add(coord, value.strip())
            if MOJIBAKE.search(value) or NONPRINTING.search(value):
                issue("mojibake", "warn").add(coord, norm(value)[:30])
            if value != value.strip() or "  " in value:
                issue("whitespace", "info").add(coord, repr(value)[:30])

    columns = []
    for i in range(width):
        cells = [(header_index + 2 + r, row[i] if i < len(row) else None) for r, row in enumerate(data)]
        live = [(n, v) for n, v in cells if not is_blank(v)]
        values = [v for _, v in live]
        letter = get_column_letter(i + 1)
        if not values and cells:
            issue("empty_columns", "info").add(f"{letter}{header_index + 1}", names[i])
        kind, confidence = classify(names[i], values, formats.get(i, ""))
        for n, v in live:
            if isinstance(v, str) and parse_number(v) is not None and kind in ("number", "percent", "currency", "mixed"):
                issue("number_as_text", "warn").add(f"{letter}{n}", norm(v)[:20])
        written = date_formats.get(i, {})
        dated = [s for n, v in live if (s := date_written(v, written.get(n)))]
        if len(set(dated)) > 1 and len(dated) / max(1, len(live)) >= 0.6:
            issue("mixed_dates", "warn").add(f"{letter}{header_index + 1}", ", ".join(sorted(set(dated))))
        by_fold: dict[str, set[str]] = defaultdict(set)
        for _, v in live:
            if isinstance(v, str):
                by_fold[norm(v).lower()].add(norm(v))
        if kind in ("text", "identifier", "boolean", "mixed") and len(by_fold) <= CATEGORICAL_MAX:
            for fold, spellings in by_fold.items():
                if len(spellings) > 1:
                    issue("casing", "info").add(f"{letter}{header_index + 1}", " / ".join(sorted(spellings)[:3]))
        samples, seen_sample = [], set()
        for v in values:
            text = norm(v)
            if text not in seen_sample:
                seen_sample.add(text)
                samples.append(text[:40])
            if len(samples) == SAMPLES:
                break
        columns.append({
            "name": names[i], "column": letter, "type": kind, "confidence": confidence,
            "missing_pct": round(100 * (len(cells) - len(live)) / len(cells), 1) if cells else 100.0,
            "unique": len({norm(v).lower() for v in values}), "samples": samples,
        })

    ordered = sorted((x.as_dict() for x in issues.values()), key=lambda d: {"fail": 0, "warn": 1, "info": 2}[d["level"]])
    return {"sheet": name, "header_row": header_index + 1, "data_rows": len(data),
            "truncated": truncated, "columns": columns, "issues": ordered}


class SheetNotFound(Exception):
    """A --sheet name the workbook does not have; profiling nothing is not a success."""


def capped(rows: Iterable, max_rows: int) -> tuple[list[list], bool]:
    """Rows up to the limit, and whether there were more: one row past it is read to tell."""
    read = [list(row) for _, row in zip(range(max_rows + 1), rows)]
    return read[:max_rows], len(read) > max_rows


def sniff_delimiter(sample: str) -> str:
    """The separator csv.Sniffer reads out of a decoded sample, comma when it cannot tell.

    A genuinely single-column file gives the sniffer nothing to measure and it raises;
    comma is the right answer there, since one column reads the same on any separator.
    """
    if len(sample) == CSV_SNIFF_BYTES:
        sample = sample.rpartition("\n")[0] or sample  # a half-read final row is not a row
    try:
        return csv.Sniffer().sniff(sample, delimiters=CSV_DELIMITERS).delimiter
    except csv.Error:
        return ","


def read_csv(path: Path, max_rows: int, delimiter: str | None) -> tuple[list[list], bool, str, str]:
    """Rows, whether they were capped, the encoding that read them, and the separator used.

    Decoding strictly and falling back beats `errors="replace"`, which turns a
    Windows-1252 export into text full of replacement characters and calls it a success.
    latin-1 maps every byte, so one candidate always works. The sniff runs on the decoded
    sample rather than the bytes, so it reads the same characters the rows will be cut on.
    """
    for encoding in CSV_ENCODINGS:
        try:
            with path.open(newline="", encoding=encoding) as fh:
                used = delimiter or sniff_delimiter(fh.read(CSV_SNIFF_BYTES))
                fh.seek(0)
                rows, truncated = capped(csv.reader(fh, delimiter=used), max_rows)
        except UnicodeDecodeError:
            continue
        return rows, truncated, encoding, used
    raise UnicodeDecodeError(CSV_ENCODINGS[-1], b"", 0, 1, "no candidate encoding read this file")


def read_sheet(ws, max_rows: int) -> SheetRead:
    """Rows up to the limit, whether there were more, what each column's format declares,
    and the format each date cell displays in, by column and sheet row.

    Read through cell objects rather than `values_only`: a percent- or currency-formatted
    column holds ordinary floats and a date cell hands back the same `datetime` whatever
    it displays, so dropping the format drops the column's type and its date mixes with it.
    """
    rows: list[list] = []
    kinds: dict[int, Counter] = defaultdict(Counter)
    dates: dict[int, dict[int, str]] = defaultdict(dict)
    truncated = False
    for row in ws.iter_rows():
        if len(rows) == max_rows:
            truncated = True
            break
        values = []
        for i, cell in enumerate(row):
            values.append(cell.value)
            if isinstance(cell.value, (datetime, date)):
                dates[i][cell.row] = cell.number_format
            elif isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                kinds[i][format_kind(cell.number_format)] += 1
        rows.append(values)
    formats = {}
    for i, seen in kinds.items():
        kind, count = seen.most_common(1)[0]
        if kind and count / sum(seen.values()) >= 0.6:  # a stray General cell does not unset a column
            formats[i] = kind
    return rows, truncated, formats, dict(dates)


def read_rows(path: Path, only_sheet: str | None, max_rows: int, delimiter: str | None) -> tuple[dict[str, SheetRead], str | None, str | None]:
    if path.suffix.lower() == ".csv":
        rows, truncated, encoding, used = read_csv(path, max_rows, delimiter)
        return {path.stem: (rows, truncated, {}, {})}, encoding, used
    wb = load_workbook(path, data_only=True, read_only=True)
    titles = [ws.title for ws in wb.worksheets]
    if only_sheet is not None and only_sheet not in titles:
        wb.close()
        raise SheetNotFound(f"--sheet {only_sheet!r} is not in this workbook; it has {titles}")
    sheets: dict[str, SheetRead] = {}
    for ws in wb.worksheets:
        if only_sheet and ws.title != only_sheet:
            continue
        sheets[ws.title] = read_sheet(ws, max_rows)
    wb.close()
    return sheets, None, None


USAGE = "usage: profile.py <file.xlsx|.csv> [--sheet NAME] [--max-rows N] [--delimiter X]"
TAKES_VALUE = {"--sheet": True, "--max-rows": True, "--delimiter": True}


def fail(message: str) -> None:
    print(json.dumps({"status": "error", "message": message}))
    sys.exit(1)


def parse_args(argv: list[str]) -> tuple[list[str], dict[str, str | bool]]:
    """Positional arguments and flags, refusing any token TAKES_VALUE does not name.

    A misspelt flag used to fall through unread, so `--shet Data` profiled every
    sheet and reported success; a value flag at the end of argv raised IndexError.
    """
    args: list[str] = []
    flags: dict[str, str | bool] = {}
    i = 0
    while i < len(argv):
        token = argv[i]
        if not token.startswith("--"):
            args.append(token)
            i += 1
        elif token not in TAKES_VALUE:
            fail(f"unknown option: {token}; {USAGE}")
        elif token in flags:
            fail(f"repeated option: {token}; {USAGE}")
        elif not TAKES_VALUE[token]:
            flags[token] = True
            i += 1
        elif i + 1 >= len(argv) or argv[i + 1].startswith("--"):
            fail(f"{token} requires a value; {USAGE}")
        else:
            flags[token] = argv[i + 1]
            i += 2
    return args, flags


def main(argv: list[str]) -> None:
    if "-h" in argv or "--help" in argv:
        print(__doc__.strip())
        sys.exit(0)
    args, flags = parse_args(argv)
    if not args:
        fail(USAGE)
    sheet = flags.get("--sheet")
    try:
        max_rows = int(flags.get("--max-rows", 5000))
    except ValueError:
        fail(f"--max-rows must be a whole number of rows; {USAGE}")
    if max_rows < 1:
        # 0 reported success over no rows at all, and a negative one never met the
        # equality stop on the workbook path, so it read every row with the cap ignored.
        fail(f"--max-rows must be at least 1 row, not {max_rows}; {USAGE}")
    delimiter = flags.get("--delimiter")
    if delimiter is not None and len(delimiter) != 1:
        fail(f"--delimiter takes a single character, not {delimiter!r}; {USAGE}")
    path = Path(args[0]).expanduser().resolve()
    try:
        sheets, encoding, separator = read_rows(path, sheet, max_rows, delimiter)
    except SheetNotFound as exc:
        print(json.dumps({"status": "error", "file": str(path), "message": str(exc)}))
        sys.exit(1)
    except Exception as exc:  # a workbook this script cannot open is a report, not a traceback
        print(json.dumps({"status": "error", "file": str(path), "message": f"{type(exc).__name__}: {exc}"}))
        sys.exit(1)
    report = {"status": "success", "file": str(path)}
    if encoding:  # a workbook has no one text encoding; only a .csv was decoded to get here
        report["encoding"] = encoding
    if separator:
        report["delimiter"] = separator
    report["sheets"] = [profile_sheet(name, rows, cut, fmts, dates) for name, (rows, cut, fmts, dates) in sheets.items()]
    print(json.dumps(report, indent=2, default=str))


if __name__ == "__main__":
    main(sys.argv[1:])
